"""Writes every validation the agent performs for a case into the test-case workbook (a new tab; the existing tabs are not touched).
The rows come from the latest data reports of the case, so the list is exactly what the agent really checks."""
from __future__ import annotations

import csv
import re
import shutil
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .bugs import fingerprint
from .config import Config
from .store import Store

KINDS = [("master_data", "Master"), ("house_data", "House"), ("manifest_data", "Manifest"), ("invoice_data", "Invoice"),
         ("packing_data", "Packing List"), ("debitnote_data", "Debit Note"), ("shipperdoc_data", "Shipper Document")]
HEAD = ["TC ID", "Area", "Field / check", "How it is validated", "Expected (from the source document)", "Last result",
        "Baseline", "Bug sheet / note"]
TAB = "Validations REG-001 & REG-002"


def latest(reports: Path, prefix: str, case_id: str) -> list[dict]:
    for f in sorted(reports.glob(f"{prefix}_*.csv"), reverse=True):
        with f.open(newline="") as fh:
            rows = [r for r in csv.DictReader(fh) if r["case_id"] == case_id]
        if rows:
            return rows
    return []


def how(result: str, expected: str) -> str:
    if result in ("PASS", "FAIL", "REVIEW", "KNOWN"):
        return "Value in Nexus must equal the value printed on the source document"
    if result == "INFO":
        return "Nexus master data or empty with nothing to compare: shown, not compared"
    return "Not readable from this sample: shown only"


def mapping_rows(reports: Path, case_id: str) -> list[dict]:
    for f in sorted(reports.glob("validate_*.csv"), reverse=True):
        with f.open(newline="") as fh:
            for r in csv.DictReader(fh):
                if r["case_id"] == case_id and r["details"]:
                    return [{"part": p} for p in r["details"].split(" | ") if re.match(r"^[A-Z ]+: ", p)]
    return []


def bug_note(store: Store, fp: str) -> tuple[str, str]:
    rec = store.bug_get(fp)
    if not rec:
        return "OPEN - not yet decided", ""
    if rec["status"] == "NOT_A_BUG":
        return "Accepted (not a bug)", rec.get("note") or ""
    m = re.search(r"!A(\d+):", rec.get("note") or "")
    return ("Known issue (logged)" if rec["status"] in ("LOGGED", "APPROVED") else "OPEN - not yet decided",
            f"UAT bug sheet row {m.group(1)}" if m else (rec.get("note") or ""))


def write(cfg: Config, case_ids: list[str], workbook: Path | None = None) -> tuple[Path, Path, int]:
    store = Store(cfg.state_dir / "agent.db")
    reports = cfg.runs_dir / "_reports"
    wb_path = workbook or sorted(cfg.path("knowledge_base/testcases").glob("Air_Single_House*.xlsx"))[-1]
    backup = wb_path.with_name(wb_path.stem + f".backup_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    shutil.copy2(wb_path, backup)
    wb = openpyxl.load_workbook(wb_path)
    if TAB in wb.sheetnames:
        del wb[TAB]
    ws = wb.create_sheet(TAB)
    ws.append(HEAD)
    n = 0
    for cid in case_ids:
        for r in mapping_rows(reports, cid):
            part = r["part"]
            area, status = part.split(": ", 1)
            ok = status.startswith("PASS")
            fp = fingerprint(cid, "Mapping", area, "", part)
            if ok:
                base, note = "Must pass", ""
            elif status.startswith("REVIEW"):
                base, note = "Accepted (alternate type)", "Nexus shows it as another type; accepted by the tester (type_alternates in settings)"
            else:
                base, note = bug_note(store, fp)
            ws.append([cid, "Mapping", area.title(), "Document of this type must sit under the parent the test case expects",
                       re.sub(r"^PASS \(|\)$", "", status) if ok else "(see the Expected Mapping column of the test case)",
                       status.split(" ")[0], base, note])
            n += 1
        for prefix, title in KINDS:
            for r in latest(reports, prefix, cid):
                if not r["field"]:
                    continue
                res, exp, shown = r["result"], " ".join(r["expected_from_document"].split()), " ".join(r["shown_in_nexus"].split())
                if res in ("PASS",):
                    base, note = "Must pass", ""
                elif res in ("FAIL", "REVIEW"):
                    base, note = bug_note(store, fingerprint(cid, title, r["field"], exp, shown))
                else:
                    base, note = "Shown only", r["note"]
                ws.append([cid, title, r["field"].split(" · ", 1)[-1], how(res, exp), exp[:300] or "-", res, base, note])
                n += 1
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEFGH", (9, 16, 34, 52, 46, 14, 26, 36)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes, ws.auto_filter.ref = "A2", ws.dimensions
    wb.save(wb_path)
    return wb_path, backup, n
